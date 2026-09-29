import csv

from legal_services.env_paths import usuarios_csv_path, usuarios_xlsx_path


def _load_from_csv() -> list[dict]:
    path = usuarios_csv_path()
    if not path.exists():
        return []
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _load_from_xlsx() -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    path = usuarios_xlsx_path()
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(h).strip().lower() for h in rows[0]]
    users = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        users.append({headers[i]: (row[i] or "") for i in range(len(headers))})
    return users


def load_users() -> list[dict]:
    users = _load_from_xlsx()
    if users:
        return users
    return _load_from_csv()


def _normalize_despacho(name: str) -> str:
    text = (name or "").strip()
    for old in ("BeChange Legal", "BeChange", "bechange legal", "bechange"):
        text = text.replace(old, "JH consulting")
    return text


def despacho_name_for_user(username: str, fallback: str = "") -> str:
    user = (username or "").strip()
    for row in load_users():
        if str(row.get("usuario", "")).strip() == user:
            return _normalize_despacho(str(row.get("nombre_despacho", user)))
    return _normalize_despacho(fallback or user)


def authenticate(username: str, password: str) -> dict | None:
    user = (username or "").strip()
    pwd = (password or "").strip()
    for row in load_users():
        if str(row.get("usuario", "")).strip() == user and str(row.get("clave", "")).strip() == pwd:
            return {
                "usuario": user,
                "tenant_id": str(row.get("tenant_id", user)).strip().lower(),
                "nombre_despacho": despacho_name_for_user(user, str(row.get("nombre_despacho", user))),
            }
    return None
