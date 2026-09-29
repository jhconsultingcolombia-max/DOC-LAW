"""Citación enviada al destinatario (contenido del correo / págs. 2-3 del PDF)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from legal_services.notificacion_constancia_docx import _CENTER_HEADERS, _add_center_line


def generate_citacion_docx(record: dict, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("NOTIFICACIÓN ELECTRÓNICA ENVIADA")
    tr.bold = True
    tr.font.size = Pt(13)
    tr.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(record.get("asunto") or "").italic = True

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run("Para: ").bold = True
    meta.add_run(record.get("para") or "")

    cuerpo = (record.get("cuerpo") or "").replace("\r\n", "\n")
    doc.add_paragraph()
    for line in cuerpo.split("\n"):
        text = line.strip()
        if not text:
            doc.add_paragraph()
            continue
        low = text.lower()
        if any(h in low for h in _CENTER_HEADERS):
            bold = "notificación personal" in low
            _add_center_line(doc, text, bold=bold, size=12 if bold else 11)
            continue
        doc.add_paragraph(text)

    doc.save(dest)
    return dest
