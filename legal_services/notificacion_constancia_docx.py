"""Constancia descargable de notificación (formato citación, contenido del mensaje enviado)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from legal_services.time_co import format_iso_colombia, now_colombia_str

_CENTER_HEADERS = (
    "república de colombia",
    "rama judicial",
    "consejo superior",
    "notificación personal",
)


def _add_center_line(doc: Document, text: str, *, bold: bool = False, size: int = 11) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def generate_constancia_docx(record: dict, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("CONSTANCIA DE NOTIFICACIÓN ELECTRÓNICA")
    tr.bold = True
    tr.font.size = Pt(14)
    tr.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run((record.get("asunto") or "")).italic = True

    doc.add_paragraph()
    meta = [
        ("Destinatario (Para)", record.get("para") or "—"),
        ("Enviado", format_iso_colombia(record.get("enviado_en") or record.get("creado"))),
        ("Apertura registrada", format_iso_colombia(record.get("abierto_en"))),
        ("Estado envío", record.get("estado") or "—"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))

    doc.add_paragraph()
    sep = doc.add_paragraph()
    sep.add_run("Contenido de la notificación").bold = True

    cuerpo = (record.get("cuerpo") or "").replace("\r\n", "\n")
    for line in cuerpo.split("\n"):
        text = line.strip()
        if not text:
            doc.add_paragraph()
            continue
        low = text.lower()
        if any(h in low for h in _CENTER_HEADERS):
            bold = "notificación personal" in low
            _add_center_line(doc, text, bold=bold, size=12 if bold else 11)
            continue
        if text.startswith("Link "):
            p = doc.add_paragraph()
            p.add_run(text)
            continue
        p = doc.add_paragraph()
        p.add_run(text)

    if record.get("adjuntos"):
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run("Archivos adjuntos al correo:").bold = True
        for adj in record["adjuntos"]:
            doc.add_paragraph(adj.get("nombre") or "—", style="List Bullet")

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run(f"Generado {now_colombia_str()} (hora Colombia)").font.size = Pt(9)

    doc.save(dest)
    return dest
