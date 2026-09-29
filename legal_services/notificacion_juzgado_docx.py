"""Memorial / constancia al juzgado (pág. 1 del PDF de referencia)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

PH_JUZGADO = {
    "juzgado_destinatario": "Ingresar nombre del juzgado (ej. Juzgado 12 Civil Municipal de Medellín)",
    "demandante": "Ingresar nombre del demandante",
    "demandado": "Ingresar nombre del demandado",
    "radicado": "Ingresar radicado del proceso",
    "cuerpo_solicitud": (
        "Actuando como apoderado de la parte demandante en el proceso de la referencia aporto "
        "constancia de recibido por ambos demandados de la notificación electrónica, solicito al "
        "despacho dar trámite a la misma."
    ),
    "firmante_nombre": "Ingresar nombre del apoderado",
    "firmante_cedula": "Ingresar cédula",
    "firmante_tp": "Ingresar tarjeta profesional",
}


def default_juzgado_campos(parsed: dict | None = None) -> dict:
    p = parsed or {}
    out = dict(PH_JUZGADO)
    for key in out:
        val = (p.get(key) or "").strip()
        if val and not val.startswith("Ingresar"):
            out[key] = val
    if p.get("juzgado_nombre") and out["juzgado_destinatario"].startswith("Ingresar"):
        out["juzgado_destinatario"] = p["juzgado_nombre"]
    return out


def generate_juzgado_docx(campos: dict, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    c = {**PH_JUZGADO, **(campos or {})}
    doc = Document()

    doc.add_paragraph("Señores")
    doc.add_paragraph(c.get("juzgado_destinatario") or "")
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.add_run("ASUNTO. ").bold = True
    p.add_run("Constancia envío notificación electrónica - personal")

    p = doc.add_paragraph()
    p.add_run("DEMANDANTE: ").bold = True
    p.add_run(c.get("demandante") or "")

    p = doc.add_paragraph()
    p.add_run("DEMANDANDO: ").bold = True
    p.add_run(c.get("demandado") or "")

    p = doc.add_paragraph()
    p.add_run("RADICADO ").bold = True
    rad = (c.get("radicado") or "").strip()
    p.add_run(rad if rad.endswith(".") else f"{rad}.")

    doc.add_paragraph()
    doc.add_paragraph(c.get("cuerpo_solicitud") or "")
    doc.add_paragraph()
    doc.add_paragraph("Atentamente,")
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph(c.get("firmante_nombre") or "")
    doc.add_paragraph(f"Cédula de Ciudadanía No. {c.get('firmante_cedula') or ''}")
    doc.add_paragraph(f"T.P. {c.get('firmante_tp') or ''} del Consejo Superior de la Judicatura.")

    for para in doc.paragraphs:
        for run in para.runs:
            run.font.size = Pt(11)

    doc.save(dest)
    return dest
