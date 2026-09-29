from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x23, 0x32)


def _add_bullets(doc: Document, items: list[str]) -> None:
    for item in items or []:
        doc.add_paragraph(item, style="List Bullet")


def generate_analysis_docx(caso: dict, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("INFORME DE ANÁLISIS PRELIMINAR")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("JH consulting · DOC_LAW · Colombia").italic = True

    doc.add_paragraph()

    meta = [
        ("Expediente", caso.get("case_id", "—")),
        ("Título", caso.get("titulo", "—")),
        ("Cliente", caso.get("cliente_nombre", "—")),
        ("Empresa", caso.get("empresa_nombre") or "—"),
        ("Rama del derecho", caso.get("rama_nombre") or caso.get("rama_id") or "—"),
        ("Abogado responsable", caso.get("abogado_responsable", "—")),
        ("Estado", caso.get("estado", "—")),
        ("Fecha informe", datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))

    _add_heading(doc, "1. Relato de hechos", 2)
    doc.add_paragraph(caso.get("hechos") or "No registrado.")

    if caso.get("entrevista"):
        _add_heading(doc, "2. Notas de entrevista", 2)
        doc.add_paragraph(caso.get("entrevista"))

    _add_heading(doc, "3. Partes del proceso", 2)
    partes = caso.get("partes") or []
    if partes:
        for p in partes:
            doc.add_paragraph(
                f"{p.get('tipo', 'parte').title()}: {p.get('nombre', '—')} "
                f"({p.get('documento') or 'sin documento'})"
            )
    else:
        doc.add_paragraph("Sin partes registradas.")

    if caso.get("procesos_radicados"):
        _add_heading(doc, "4. Procesos radicados", 2)
        doc.add_paragraph(caso.get("procesos_radicados"))

    viab = caso.get("viabilidad") or {}
    _add_heading(doc, "5. Viabilidad preliminar", 2)
    doc.add_paragraph(f"Resultado: {viab.get('viabilidad_preliminar', '—')}")
    if viab.get("alertas"):
        _add_bullets(doc, viab["alertas"])

    analysis = caso.get("analisis_legal") or {}
    _add_heading(doc, "6. Análisis jurídico preliminar", 2)

    doc.add_paragraph("A favor del cliente:").runs[0].bold = True
    _add_bullets(doc, analysis.get("analisis_a_favor_cliente") or [])

    doc.add_paragraph("En contra / riesgos:").runs[0].bold = True
    _add_bullets(doc, analysis.get("analisis_en_contra") or [])

    doc.add_paragraph("Pasos sugeridos:").runs[0].bold = True
    _add_bullets(doc, analysis.get("pasos_sugeridos") or [])

    ai = caso.get("analisis_ia") or {}
    if ai.get("ia_disponible"):
        _add_heading(doc, "7. Análisis con IA — Checklist documental", 2)
        for item in ai.get("checklist_documentos") or []:
            doc.add_paragraph(
                f"[{item.get('estado', '—').upper()}] {item.get('item', '—')} — {item.get('evidencia', '')}"
            )

        _add_heading(doc, "8. Hallazgos en documentos", 2)
        _add_bullets(doc, ai.get("hallazgos_documentos") or [])

        _add_heading(doc, "9. Marco legal Colombia", 2)
        for norm in ai.get("marco_legal_colombia") or []:
            doc.add_paragraph(
                f"{norm.get('norma', '—')} ({norm.get('articulo_o_referencia', '')}): "
                f"{norm.get('aplicacion_al_caso', '')}"
            )

        _add_heading(doc, "10. Reglamento / políticas empresa", 2)
        for reg in ai.get("reglamento_empresa") or []:
            doc.add_paragraph(
                f"{reg.get('documento', '—')} — {reg.get('clausula_o_politica', '')}: {reg.get('aplicacion', '')}"
            )

        _add_heading(doc, "11. Acciones recomendadas (estatuto / ley)", 2)
        for act in ai.get("acciones_recomendadas") or []:
            doc.add_paragraph(
                f"Prioridad {act.get('prioridad', '—')}: {act.get('accion', '')} — Fundamento: {act.get('fundamento', '')}"
            )

        if ai.get("resumen_ejecutivo"):
            _add_heading(doc, "12. Resumen ejecutivo IA", 2)
            doc.add_paragraph(ai.get("resumen_ejecutivo"))

    _add_heading(doc, "13. Borrador de comunicación al cliente", 2)
    doc.add_paragraph(caso.get("borrador_comunicacion") or "—")

    _add_heading(doc, "14. Fuentes oficiales de consulta", 2)
    fuentes = analysis.get("fuentes_oficiales") or {}
    _add_bullets(doc, [f"{k}: {v}" for k, v in fuentes.items()])

    disclaimer = doc.add_paragraph()
    disclaimer.add_run(
        "\n\nAVISO: Este informe es un borrador preliminar generado por JH consulting (DOC_LAW). "
        "No constituye opinión legal definitiva. Requiere revisión, validación y firma del abogado responsable."
    ).italic = True

    doc.save(str(dest))
    return dest
